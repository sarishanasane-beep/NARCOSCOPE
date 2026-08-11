"""Text extraction services.

Separate extractors for PDF, DOCX and TXT. Each returns a list of
``(page_number, raw_text)`` tuples. Failures raise ``ExtractionError``
and are caught by the processing pipeline so the application never
crashes.
"""

from __future__ import annotations

import logging
from abc import ABC, abstractmethod
from pathlib import Path

from app.utils.errors import ExtractionError

logger = logging.getLogger(__name__)


class BaseTextExtractor(ABC):
    name = "base"

    @abstractmethod
    def extract(self, path: Path) -> list[tuple[int, str]]:
        """Return ``[(page_number, text), ...]``."""


class PDFTextExtractor(BaseTextExtractor):
    """PyMuPDF (fitz) page-by-page extraction."""

    name = "pdf_pymupdf"

    def extract(self, path: Path) -> list[tuple[int, str]]:
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:  # pragma: no cover
            raise ExtractionError(
                "Text extraction failed. Please review the original document.",
                technical=f"PyMuPDF not installed: {exc}",
            )
        try:
            pages: list[tuple[int, str]] = []
            with fitz.open(str(path)) as pdf:
                for index, page in enumerate(pdf, start=1):
                    pages.append((index, page.get_text()))
            return pages
        except Exception as exc:  # noqa: BLE001 — report, don't crash
            logger.exception("pdf extraction failed path=%s", path)
            raise ExtractionError(technical=f"pdf extraction error: {exc}") from exc


class DOCXTextExtractor(BaseTextExtractor):
    """python-docx paragraph + table extraction."""

    name = "docx_python_docx"

    def extract(self, path: Path) -> list[tuple[int, str]]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            raise ExtractionError(
                "Text extraction failed. Please review the original document.",
                technical=f"python-docx not installed: {exc}",
            )
        try:
            doc = DocxDocument(str(path))
            chunks: list[str] = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    chunks.append(paragraph.text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        chunks.append(" | ".join(cells))
            return [(1, "\n".join(chunks))]
        except Exception as exc:  # noqa: BLE001
            logger.exception("docx extraction failed path=%s", path)
            raise ExtractionError(technical=f"docx extraction error: {exc}") from exc


class TXTTextExtractor(BaseTextExtractor):
    """Plain text with encoding handling."""

    name = "txt_encoding_safe"

    def extract(self, path: Path) -> list[tuple[int, str]]:
        try:
            data = path.read_bytes()
        except OSError as exc:
            raise ExtractionError(technical=f"txt read error: {exc}") from exc
        for encoding in ("utf-8-sig", "utf-16", "latin-1"):
            try:
                return [(1, data.decode(encoding))]
            except UnicodeDecodeError:
                continue
        raise ExtractionError(technical="txt decoding failed for all encodings")


def get_extractor(file_type: str) -> BaseTextExtractor:
    extractors = {
        "pdf": PDFTextExtractor(),
        "docx": DOCXTextExtractor(),
        "txt": TXTTextExtractor(),
    }
    extractor = extractors.get(file_type)
    if extractor is None:
        raise ExtractionError(technical=f"no extractor for file_type={file_type}")
    return extractor
